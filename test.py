from docx import Document
from docx.shared import Pt
import os

# โหลดไฟล์ Word ที่เพิ่งอัปเดตไปเมื่อกี้
file_name = "AgriDirect_Report_G3_Updated.docx"
if not os.path.exists(file_name):
    # Fallback เผื่อหาไฟล์ไม่เจอ ให้ใช้ไฟล์ตั้งต้น
    file_name = "AgriDirect_Report_G3.docx"

doc = Document(file_name)

# เพิ่มหัวข้อที่ 4 ต่อจาก 3 Priority Actions เดิม
doc.add_heading('4. สรุปและการดำเนินงานขั้นต่อไป (จากสไลด์ผู้บริหาร)', level=2)

doc.add_heading('4.1 การประเมินเป้าหมาย', level=3)
# เปลี่ยนมาใช้พารากราฟปกติ แล้วใส่ '•' เอาเอง
p1 = doc.add_paragraph()
p1.add_run('• เป้าหมายยอดต่อบิล +25%: ').bold = True
p1.add_run('มีความเป็นไปได้และสามารถทำได้จริงผ่านกลยุทธ์การจัดเซ็ตสินค้า (Product Bundling)')

p2 = doc.add_paragraph()
p2.add_run('• อัตราการซื้อซ้ำ (Repeat Rate): ').bold = True
p2.add_run('ตัวเลขจากการคำนวณการซื้อซ้ำจริงอยู่ที่ 78.1% ซึ่งทะลุเป้าหมายเดิมที่ตั้งไว้ (32%) ไปแล้วเรียบร้อย')

doc.add_heading('4.2 สิ่งที่ต้องทำทันที (Immediate Actions)', level=3)
p3 = doc.add_paragraph()
p3.add_run('• ทดลองย้ายงบโฆษณาไปที่ช่องทาง LINE OA เป็นระยะเวลา 4–6 สัปดาห์ เพื่อดูผลลัพธ์ที่แท้จริง')

p4 = doc.add_paragraph()
p4.add_run('• เปิดตัวเซ็ตสินค้า (Bundle) ข้ามหมวดหมู่ พร้อมกับทดสอบการปรับค่าจัดส่งแยกโซน')

doc.add_heading('4.3 ข้อจำกัดของข้อมูล (Data Limitations)', level=3)
p5 = doc.add_paragraph()
p5.add_run('• ต้นทุนสินค้า (COGS) ยังคงอ้างอิงจากตัวเลขสมมติฐานที่ 55% (ต้องใช้ต้นทุนจริงเพื่อความแม่นยำของกำไร)')

p6 = doc.add_paragraph()
p6.add_run('• ผลประเมินประสิทธิภาพช่องทางโฆษณา ใช้การจับคู่รายได้แบบประมาณการ (Proxy attribution) แทน')

# เซฟเป็นไฟล์ Final
output_path = "AgriDirect_Report_G3_Final.docx"
doc.save(output_path)
print(f"✅ อัปเดตข้อมูลสรุปลงไฟล์ Word สำเร็จ! บันทึกเป็น: {output_path}")